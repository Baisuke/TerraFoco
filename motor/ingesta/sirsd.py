# -*- coding: utf-8 -*-
"""
Carga de la ejecución del programa SIRSD-S: el eje de inversión.

Esta es la mitad del proyecto que no viene de un satélite. Sin ella no hay
pregunta que responder: el índice de pertinencia compara el estado del suelo
con la plata que llegó a cada comuna.

Acepta dos calidades de fuente y lo declara en cada fila:

  entrega directa       INDAP entrega el detalle por comuna. Trae superficie
                        bonificada -> precision_dato='completa'.
  transparencia activa  Planillas públicas mensuales. A veces sin superficie
                        -> precision_dato='reducida'. Sirve igual: el índice se
                        arma con monto y beneficiarios, y la interfaz lo dice.

Lo que NO hace es descartar filas en silencio. Todo lo que no se puede asociar
a una comuna queda en `programas.ejecucion_rechazada` con el motivo. Un
programa que pierde el 8% de los registros sin avisar produce un mapa que
miente, y nadie se entera.

Uso:
    python -m ingesta.sirsd --archivo planilla.csv --anio 2024 --simular
    python -m ingesta.sirsd --archivo planilla.csv --anio 2024 --fuente "transparencia activa"
"""
import argparse
import csv
import io
import json
import os
import re
import sys
import unicodedata

import config
from bd import Bitacora, conexion, obtener_region

FUENTES = ("entrega directa", "transparencia activa", "SAIP")

# Nombres de columna que se han visto en las planillas. Se busca por
# coincidencia flexible porque cada archivo las rotula distinto.
ALIAS = {
    "comuna":       ("comuna", "nombre comuna", "nom comuna", "des comuna"),
    "superficie":   ("superficie", "superficie ha", "sup bonificada",
                     "hectareas", "has", "superficie bonificada"),
    "monto":        ("monto", "monto pesos", "bonificacion", "monto bonificado",
                     "incentivo", "monto total"),
    "beneficiarios": ("beneficiarios", "n beneficiarios", "usuarios",
                      "cantidad beneficiarios", "n usuarios"),
    "practica":     ("practica", "subprograma", "tipo practica", "labor"),
}


def normalizar(texto):
    """Sin tildes, sin puntuación, en minúsculas y con espacios colapsados.

    'Ñuñoa' y 'NUNOA' deben caer en la misma clave, o se pierden filas por
    diferencias de tipeo que no significan nada.
    """
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^a-zA-Z0-9\s]", " ", t)
    return " ".join(t.lower().split())


def a_numero(valor):
    """Convierte '1.234,50', '$ 1.234', '1234.5' o '' a float o None.

    Las planillas chilenas usan punto de miles y coma decimal; un float() a
    secas convierte 1.234 en 1,234 y subestima el monto por mil.
    """
    if valor is None:
        return None
    s = str(valor).strip()
    if not s or s in ("-", "s/i", "S/I", "N/A"):
        return None
    s = re.sub(r"[^\d,.\-]", "", s)
    if not s:
        return None

    if "," in s and "." in s:
        # 1.234,50 — punto de miles, coma decimal
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        # 1234,50
        s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        # 12.500.000 — solo puntos, en grupos exactos de tres cifras. Se
        # interpretan como miles, que es como vienen las planillas chilenas.
        # Sin esta rama, float('12.500.000') lanza ValueError y el monto se
        # perdia devolviendo None: la fila entraba con plata en cero y bajaba
        # la intensidad de intervencion de la comuna sin que nada fallara.
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def mapear_columnas(cabeceras):
    """Asocia cada campo lógico con la columna real del archivo."""
    norm = {normalizar(h): h for h in cabeceras}
    mapa = {}
    for campo, alias in ALIAS.items():
        for a in alias:
            if a in norm:
                mapa[campo] = norm[a]
                break
        else:
            # segunda pasada, por contención: 'monto bonificado 2024'
            for clave, original in norm.items():
                if any(a in clave for a in alias):
                    mapa[campo] = original
                    break
    return mapa


def leer_filas(ruta):
    """Lee CSV o XLSX y devuelve (cabeceras, filas como dict)."""
    if not os.path.exists(ruta):
        raise SystemExit("No existe el archivo: %s" % ruta)

    if ruta.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("Falta openpyxl. Instalar con: pip install openpyxl")
        hoja = load_workbook(ruta, read_only=True, data_only=True).active
        it = hoja.iter_rows(values_only=True)
        cabeceras = [str(c) if c is not None else "" for c in next(it)]
        return cabeceras, [dict(zip(cabeceras, f)) for f in it]

    # El separador y la codificación varían entre planillas; se detectan.
    with io.open(ruta, "rb") as f:
        crudo = f.read()
    for cod in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = crudo.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4096]
    sep = ";" if muestra.count(";") > muestra.count(",") else ","
    lector = csv.DictReader(io.StringIO(texto), delimiter=sep)
    return lector.fieldnames or [], list(lector)


def catalogo_comunas(con, region):
    """Nombre normalizado -> id_comuna, para la región activa."""
    filas = con.execute(
        "SELECT id_comuna, nombre FROM territorio.comuna WHERE id_region = %s",
        (region,),
    ).fetchall()
    return {normalizar(n): i for i, n in filas}


def id_programa(con, organismo):
    fila = con.execute(
        "SELECT id_programa FROM programas.programa "
        "WHERE codigo = 'SIRSD-S' AND organismo = %s",
        (organismo,),
    ).fetchone()
    if fila is None:
        raise SystemExit(
            "No existe el programa SIRSD-S para el organismo %r. "
            "Organismos válidos: INDAP, SAG." % organismo
        )
    return fila[0]


def catalogo_practicas(con):
    filas = con.execute(
        "SELECT id_practica, nombre FROM programas.practica"
    ).fetchall()
    return {normalizar(n): i for i, n in filas}


def emparejar_practica(texto, practicas):
    """Coincidencia laxa: las planillas abrevian los nombres de las prácticas."""
    if not texto:
        return None
    clave = normalizar(texto)
    if clave in practicas:
        return practicas[clave]
    for nombre, id_p in practicas.items():
        if clave in nombre or nombre in clave:
            return id_p
    palabras = set(clave.split())
    for nombre, id_p in practicas.items():
        if len(palabras & set(nombre.split())) >= 2:
            return id_p
    return None


SQL_EJECUCION = """
INSERT INTO programas.ejecucion
    (id_programa, id_comuna, id_practica, anio, superficie_ha, monto_pesos,
     beneficiarios, precision_dato, fuente)
VALUES (%(id_programa)s, %(id_comuna)s, %(id_practica)s, %(anio)s,
        %(superficie_ha)s, %(monto_pesos)s, %(beneficiarios)s,
        %(precision_dato)s, %(fuente)s)
"""

SQL_RECHAZO = """
INSERT INTO programas.ejecucion_rechazada (linea_origen, motivo, archivo)
VALUES (%s, %s, %s)
"""


def procesar(filas, mapa, comunas, practicas, anio, fuente, id_prog):
    """Separa lo cargable de lo rechazado. No descarta nada en silencio."""
    listos, rechazos = [], []

    for fila in filas:
        crudo = {str(k): (None if v is None else str(v)) for k, v in fila.items()}
        nombre = fila.get(mapa.get("comuna", ""), None)

        if not nombre or not str(nombre).strip():
            rechazos.append((crudo, "sin nombre de comuna"))
            continue

        id_comuna = comunas.get(normalizar(nombre))
        if id_comuna is None:
            rechazos.append((crudo, "comuna no encontrada en la region: %r"
                             % str(nombre)[:60]))
            continue

        superficie = a_numero(fila.get(mapa.get("superficie", "")))
        monto = a_numero(fila.get(mapa.get("monto", "")))
        benef = a_numero(fila.get(mapa.get("beneficiarios", "")))

        if superficie is None and monto is None and benef is None:
            rechazos.append((crudo, "fila sin superficie, monto ni beneficiarios"))
            continue

        # La precisión no se adivina: se deduce de lo que trae la fila.
        precision = "completa" if superficie is not None else "reducida"

        listos.append({
            "id_programa": id_prog,
            "id_comuna": id_comuna,
            "id_practica": emparejar_practica(
                fila.get(mapa.get("practica", "")), practicas),
            "anio": anio,
            "superficie_ha": round(superficie, 2) if superficie is not None else None,
            "monto_pesos": round(monto) if monto is not None else None,
            "beneficiarios": int(benef) if benef is not None else None,
            "precision_dato": precision,
            "fuente": fuente,
        })

    return listos, rechazos


def resumen(listos, rechazos, total):
    print("\n[sirsd] %d filas leidas" % total)
    print("        %d cargables" % len(listos))
    print("        %d rechazadas" % len(rechazos))
    if total:
        print("        %.1f%% de perdida" % (100.0 * len(rechazos) / total))

    if rechazos:
        motivos = {}
        for _c, m in rechazos:
            clave = m.split(":")[0]
            motivos[clave] = motivos.get(clave, 0) + 1
        print("\n        motivos del rechazo:")
        for m, n in sorted(motivos.items(), key=lambda x: -x[1]):
            print("          %-45s %5d" % (m, n))

    if listos:
        porcom = {}
        for r in listos:
            d = porcom.setdefault(r["id_comuna"], {"ha": 0.0, "monto": 0})
            d["ha"] += r["superficie_ha"] or 0
            d["monto"] += r["monto_pesos"] or 0
        reducidas = sum(1 for r in listos if r["precision_dato"] == "reducida")
        print("\n        %d comunas con ejecucion" % len(porcom))
        if reducidas:
            print("        %d filas con precision reducida (sin superficie)" % reducidas)


def main():
    p = argparse.ArgumentParser(description="Carga la ejecución del SIRSD-S")
    p.add_argument("--archivo", required=True, help="planilla CSV o XLSX")
    p.add_argument("--anio", type=int, required=True)
    p.add_argument("--organismo", default="INDAP", choices=("INDAP", "SAG"))
    p.add_argument("--fuente", default="transparencia activa", choices=FUENTES)
    p.add_argument("--region", type=int, default=config.ID_REGION)
    p.add_argument("--simular", action="store_true",
                   help="analiza y muestra el resumen, pero no escribe")
    args = p.parse_args()

    cabeceras, filas = leer_filas(args.archivo)
    mapa = mapear_columnas(cabeceras)

    print("[sirsd] columnas detectadas:")
    for campo in ALIAS:
        print("        %-14s -> %s" % (campo, mapa.get(campo, "(no encontrada)")))
    if "comuna" not in mapa:
        raise SystemExit(
            "No se encontro la columna de comuna. Cabeceras del archivo:\n  %s"
            % ", ".join(cabeceras)
        )

    region = obtener_region(args.region)

    with Bitacora("Ingesta SIRSD-S %d — %s" % (args.anio, region.nombre)) as b:
        with conexion() as con:
            comunas = catalogo_comunas(con, args.region)
            practicas = catalogo_practicas(con)
            id_prog = id_programa(con, args.organismo)

            listos, rechazos = procesar(filas, mapa, comunas, practicas,
                                        args.anio, args.fuente, id_prog)
            resumen(listos, rechazos, len(filas))

            if args.simular:
                print("\n[sirsd] --simular: no se escribio nada.")
                b.simulada = True
                b.registros = 0
                return 0

            for r in listos:
                con.execute(SQL_EJECUCION, r)
            for crudo, motivo in rechazos:
                con.execute(SQL_RECHAZO,
                            (json.dumps(crudo, ensure_ascii=False), motivo,
                             os.path.basename(args.archivo)))
            b.registros = len(listos)
            print("\n[sirsd] %d filas cargadas, %d en ejecucion_rechazada"
                  % (len(listos), len(rechazos)))

    return 0


if __name__ == "__main__":
    sys.exit(main())
