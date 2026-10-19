# -*- coding: utf-8 -*-
"""
Agrega la entrega oficial de INDAP ("1- Solo Planes <año>.xlsx") por comuna y
año, para completar y validar la serie del SIRSD-S.

LAS PLANILLAS TRAEN DATOS PERSONALES: RUT, nombres, sexo, etnia, rol y nombre
del predio de beneficiarios y operadores. Esta herramienta:

  - lee SOLO columnas no personales, por nombre (COLUMNAS, abajo);
  - suma por comuna y año, y escribe únicamente esos agregados;
  - no copia las planillas ni deja nada en disco salvo el CSV de salida.

Las planillas no se versionan ni se cargan en TerraFoco. Deben quedar fuera
del repositorio y, en lo posible, fuera de carpetas que se sincronizan a la
nube.

Definición, la misma de la serie definitiva (docs/17-serie-sirsd.md):
  - ejecutado = ESTADO PM «Declaración Jurada» o «Traspasado a Tesorería»
  - año = año de FECHA DEC JURADA PM (la entrega viene por año de concurso)
  - comuna = COMUNA PREDIO, llevada a código INE

Uso:
    python herramientas/agregar_entrega_indap.py "<carpeta con las planillas>"
    python herramientas/agregar_entrega_indap.py "<carpeta>" --salida otro.csv
"""
import argparse
import glob
import json
import os
import sys
import unicodedata

import pandas as pd
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "bd", "semillas", "sirsd_complemento_oficial_2012_2025.csv")
COMUNAS = os.path.join(RAIZ, "prototipo", "datos", "comunas.geojson")

# Únicas columnas que se leen. Ninguna identifica a una persona.
COLUMNAS = ["CODIGO PM", "ESTADO PM", "FECHA DEC JURADA PM", "COMUNA PREDIO",
            "MONTO INVERSION", "MONTO INCENTIVO", "MONTO ASES FORMULACION",
            "MONTO ASES EJECUCION", "SUPAGR PREDIO", "HA REALES EJECUTADA PM"]
EJECUTADO = ("Declaración Jurada", "Traspasado a Tesorería")
# INDAP escribe algunas comunas distinto que el INE.
ALIAS = {"QUINTA TILCOCO": "QUINTA DE TILCOCO"}


def norm(t):
    t = unicodedata.normalize("NFD", str(t)).encode("ascii", "ignore").decode().upper().strip()
    return ALIAS.get(t, t)


def codigos_ine():
    with open(COMUNAS, encoding="utf-8") as f:
        feats = json.load(f)["features"]
    return {norm(x["properties"]["nombre"]): x["properties"]["id_comuna"] for x in feats}


def leer(carpeta):
    filas = []
    archivos = sorted(glob.glob(os.path.join(carpeta, "*Solo Planes*.xlsx")))
    if not archivos:
        raise SystemExit("No hay planillas «Solo Planes» en %s" % carpeta)
    for ruta in archivos:
        wb = load_workbook(ruta, read_only=True)
        it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(it)]
        faltan = [c for c in COLUMNAS if c not in cab]
        if faltan:
            raise SystemExit("%s: faltan columnas %s" % (os.path.basename(ruta), faltan))
        idx = [cab.index(c) for c in COLUMNAS]
        for r in it:
            if r and any(v is not None for v in r):
                filas.append([r[i] for i in idx])
        wb.close()
    print("[entrega] %d planillas, %d planes" % (len(archivos), len(filas)))
    return pd.DataFrame(filas, columns=COLUMNAS)


def agregar(o, codigos):
    o = o.copy()
    o["id_comuna"] = o["COMUNA PREDIO"].map(lambda n: codigos.get(norm(n)))
    sin = sorted(set(o.loc[o.id_comuna.isna(), "COMUNA PREDIO"].astype(str)))
    if sin:
        raise SystemExit("Comunas sin código INE: %s (agregarlas a ALIAS)" % sin)
    o = o[o["ESTADO PM"].isin(EJECUTADO)]
    fecha = pd.to_datetime(o["FECHA DEC JURADA PM"], errors="coerce", dayfirst=True)
    print("[entrega] %d ejecutados; %d sin fecha de declaración quedan fuera (%.1f%%)"
          % (len(o), fecha.isna().sum(), 100 * fecha.isna().mean()))
    o = o.assign(anio=fecha.dt.year)
    o = o[o.anio.between(2012, 2025)]
    for c in COLUMNAS[4:]:
        o[c] = pd.to_numeric(o[c], errors="coerce")
    return (o.groupby(["id_comuna", "anio"])
             .agg(planes_oficial=("CODIGO PM", "nunique"),
                  incentivo_oficial=("MONTO INCENTIVO", "sum"),
                  inversion_oficial=("MONTO INVERSION", "sum"),
                  ha_ejecutada_oficial=("HA REALES EJECUTADA PM", "sum"),
                  ases_formulacion=("MONTO ASES FORMULACION", "sum"),
                  ases_ejecucion=("MONTO ASES EJECUCION", "sum"),
                  supagr_predios=("SUPAGR PREDIO", "sum"))
             .reset_index()
             .astype({"id_comuna": int, "anio": int}))


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    p.add_argument("carpeta", help="carpeta con las planillas «1- Solo Planes <año>.xlsx»")
    p.add_argument("--salida", default=SALIDA)
    args = p.parse_args()

    tabla = agregar(leer(args.carpeta), codigos_ine())
    # Última defensa: la salida solo puede tener estas columnas, todas numéricas.
    assert list(tabla.columns) == ["id_comuna", "anio", "planes_oficial", "incentivo_oficial",
                                   "inversion_oficial", "ha_ejecutada_oficial",
                                   "ases_formulacion", "ases_ejecucion", "supagr_predios"]
    assert all(pd.api.types.is_numeric_dtype(t) for t in tabla.dtypes)
    os.makedirs(os.path.dirname(args.salida), exist_ok=True)
    tabla.round(2).to_csv(args.salida, sep=";", index=False)
    print("[entrega] %d celdas comuna-año -> %s" % (len(tabla), os.path.relpath(args.salida, RAIZ)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
